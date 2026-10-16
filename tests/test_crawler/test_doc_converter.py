from __future__ import annotations

import json
from pathlib import Path
from unittest.mock import patch

import docx

from scripts.crawler.config import CrawlerConfig
from scripts.crawler.doc_converter import DocConverter
from scripts.crawler.doc_detector import OLE2_MAGIC


def test_verify_docx_valid(tmp_path: Path):
    doc_path = tmp_path / "valid.docx"
    doc = docx.Document()
    doc.add_paragraph("Test Paragraph")
    doc.save(str(doc_path))

    is_valid, msg = DocConverter.verify_docx(doc_path)
    assert is_valid is True
    assert msg == ""


def test_verify_docx_invalid_zip(tmp_path: Path):
    invalid_path = tmp_path / "corrupt.docx"
    invalid_path.write_bytes(b"not a zip file at all")

    is_valid, msg = DocConverter.verify_docx(invalid_path)
    assert is_valid is False
    assert "Invalid zip package" in msg


def test_converter_metadata_update(tmp_path: Path):
    config = CrawlerConfig(output_dir=tmp_path)
    converter = DocConverter(config=config)

    # Populate documents.jsonl
    original_record = {
        "document_id": "doc_test_123",
        "title": "Mẫu đơn đăng ký",
        "source_url": "https://tuyensinh.epu.edu.vn/mau.doc",
        "canonical_url": "https://tuyensinh.epu.edu.vn/mau.doc",
        "parent_url": "https://tuyensinh.epu.edu.vn/chi-tiet-tin/123.html",
        "source_id": "epu_admissions",
        "document_type": "docx",
        "admission_year": 2026,
        "major_code": "7480201",
        "local_path": "raw/docx/test_mau.doc",
        "sha256": "original_sha256_hash",
        "file_size_bytes": 65536,
        "status": "downloaded",
        "metadata_extra": {},
    }
    with open(config.documents_file, "w", encoding="utf-8") as f:
        f.write(json.dumps(original_record, ensure_ascii=False) + "\n")

    conversions_map = {
        "test_mau.doc": {
            "original_rel_path": "raw/docx/test_mau.doc",
            "converted_rel_path": "converted/docx/test_mau.docx",
            "original_sha256": "original_sha256_hash",
            "converted_sha256": "converted_sha256_hash",
            "engine": "libreoffice",
        }
    }

    converter._update_metadata(conversions_map)

    with open(config.documents_file, "r", encoding="utf-8") as f:
        updated = json.loads(f.readline().strip())

    assert updated["document_id"] == "doc_test_123"
    assert updated["source_url"] == "https://tuyensinh.epu.edu.vn/mau.doc"
    assert updated["major_code"] == "7480201"
    assert updated["converted_path"] == "converted/docx/test_mau.docx"
    assert updated["converted_sha256"] == "converted_sha256_hash"
    assert updated["metadata_extra"]["original_doc_path"] == "raw/docx/test_mau.doc"
    assert updated["metadata_extra"]["conversion_engine"] == "libreoffice"


def test_converter_orchestration_mocked(tmp_path: Path):
    config = CrawlerConfig(output_dir=tmp_path)
    converter = DocConverter(config=config)

    # Create synthetic raw doc
    raw_doc = config.docx_dir / "sample.doc"
    raw_doc.write_bytes(OLE2_MAGIC + b"\x00" * 1024)

    # Mock _convert_single to simulate successful conversion
    def fake_convert(in_f, out_f):
        # generate valid docx
        d = docx.Document()
        d.add_paragraph("Converted from doc")
        d.save(str(out_f))
        return True, "mock_libreoffice", ""

    with patch.object(converter, "_convert_single", side_effect=fake_convert):
        res = converter.run(all_files=True, dry_run=False)

    assert res["total_scanned"] == 1
    assert res["valid_doc_count"] == 1
    assert res["converted_count"] == 1
    assert res["failed_count"] == 0

    converted_docx = config.converted_docx_dir / "sample.docx"
    assert converted_docx.exists()

    is_valid, _ = DocConverter.verify_docx(converted_docx)
    assert is_valid is True


def test_convert_via_pure_python_fallback(tmp_path: Path):
    converter = DocConverter(config=CrawlerConfig(output_dir=tmp_path))
    input_doc = tmp_path / "test_legacy.doc"

    # Construct synthetic OLE2 content with UTF-16LE text runs
    text_content = "Cộng hòa Xã hội Chủ nghĩa Việt Nam\nThông tin tuyển sinh đại học".encode("utf-16le")
    payload = OLE2_MAGIC + b"\x00" * 128 + text_content + b"\x00" * 128
    input_doc.write_bytes(payload)

    target_docx = tmp_path / "test_legacy.docx"
    success, err = converter._convert_via_pure_python(input_doc, target_docx)

    assert success is True
    assert err == ""
    assert target_docx.exists()

    is_valid, val_err = DocConverter.verify_docx(target_docx)
    assert is_valid is True, f"Validation failed: {val_err}"

    # Verify document structure
    doc = docx.Document(str(target_docx))
    assert len(doc.paragraphs) > 0


def test_convert_file_with_replace_original(tmp_path: Path):
    converter = DocConverter(config=CrawlerConfig(output_dir=tmp_path))
    input_doc = tmp_path / "replace_test.doc"
    input_doc.write_bytes(OLE2_MAGIC + b"\x00" * 512)

    def fake_convert_single(in_f, out_f):
        d = docx.Document()
        d.add_paragraph("Mock paragraph")
        d.save(str(out_f))
        return True, "mock_engine", ""

    with patch.object(converter, "_convert_single", side_effect=fake_convert_single):
        success, final_path, engine = converter.convert_file(
            input_doc=input_doc,
            replace_original=True,
        )

    assert success is True
    assert engine == "mock_engine"
    assert final_path.suffix == ".docx"
    assert final_path.exists()
    assert not input_doc.exists()  # Original .doc must be removed


def test_storage_in_stream_doc_conversion(tmp_path: Path):
    from scripts.crawler.storage import LocalStorageManager

    config = CrawlerConfig(output_dir=tmp_path)
    storage = LocalStorageManager(config=config)

    content = OLE2_MAGIC + b"\x00" * 512
    url = "https://tuyensinh.epu.edu.vn/Uploads/files/ts2026/phieu_dang_ky.doc"
    sha = "test_sha_doc_123456789"

    def fake_convert_bytes(content, filename, target_dir):
        out_p = target_dir / filename
        d = docx.Document()
        d.add_paragraph("Directly converted in pipeline")
        d.save(str(out_p))
        return True, out_p, out_p.read_bytes(), "mock_word"

    with patch("scripts.crawler.doc_converter.DocConverter.convert_doc_bytes", side_effect=fake_convert_bytes):
        rel_path, is_new = storage.save_raw_file(
            url=url,
            content=content,
            sha256_hash=sha,
            doc_type="docx",
        )

    assert is_new is True
    assert rel_path.endswith(".docx")
    assert not rel_path.endswith(".doc")
    full_path = tmp_path / rel_path
    assert full_path.exists()
    is_valid, _ = DocConverter.verify_docx(full_path)
    assert is_valid is True

